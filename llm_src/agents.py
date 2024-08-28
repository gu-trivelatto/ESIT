import os
import sys
import sqlite3

from pathlib import Path
from datetime import datetime
from CESM.core.model import Model
from CESM.core.input_parser import Parser
from CESM.core.plotter import Plotter, PlotType
from CESM.core.data_access import DAO
from openpyxl import load_workbook
from abc import ABC, abstractmethod
from llm_src.state import GraphState
from llm_src.helper import HelperFunctions
from langchain.schema import Document
from langchain.prompts import PromptTemplate
from langchain_core.output_parsers import StrOutputParser
from langchain_core.output_parsers import JsonOutputParser

# TODO standardize variable names used from the state
# TODO standardize the way the agents interact with the state

class AgentBase(ABC):
    def __init__(self, llm_models, es_models, state: GraphState, app, debug):
        self.chat_model = llm_models['chat_model']
        self.json_model = llm_models['json_model']
        self.ht_model = llm_models['ht_model']
        self.ht_json_model = llm_models['ht_json_model']
        self.state = state
        self.debug = debug
        self.app = app
        self.selected_value = None
        self.base_model = es_models['base_model']
        self.mod_model = es_models['mod_model']
        self.helper = HelperFunctions()
        
    def confirm_selection(self, selected_value):
        self.selected_value = selected_value

    @abstractmethod
    def get_prompt_template(self) -> PromptTemplate:
        pass

    def execute(self) -> GraphState:
        pass
    
class ResearchAgentBase(ABC):
    def __init__(self, llm_models, retriever, web_tool, state: GraphState, app, debug):
        self.retriever = retriever
        self.web_tool = web_tool
        self.chat_model = llm_models['chat_model']
        self.json_model = llm_models['json_model']
        self.ht_model = llm_models['ht_model']
        self.ht_json_model = llm_models['ht_json_model']
        self.state = state
        self.debug = debug
        self.app = app
        self.selected_value = None
        self.helper = HelperFunctions()
        
    def get_answer_analyzer_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an expert at summarizing a bunch of data to extract only the important bits from it.

            Given the user's QUERY and the SEARCH_RESULTS, summarize as briefly as possible the information
            searched by the user. Don't give any preamble or introduction, go directly to the summary
            of the requested information.
            
            If it helps to provide a more precise answer, you can also make use of the CONTEXT.
            
            Whenever there is a source in the data received in SEARCH_RESULTS, you MUST include the used
            sources as a topic at the end of the summarization. In case there is no source, simply ignore this
            instruction.

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            QUERY: {query} \n
            SEARCH_RESULTS: {search_results} \n
            CONTEXT: {context} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["query","search_results","context"],
        )

    @abstractmethod
    def get_prompt_template(self) -> PromptTemplate:
        pass

    def execute(self) -> GraphState:
        pass

class DateGetter(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return super().get_prompt_template()
    
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Processing')
        num_steps = self.state['num_steps']
        num_steps += 1
        
        current_date = datetime.now().strftime("%d %B %Y, %H:%M:%S")
        
        result = f'The current date and time are {current_date}'
        
        if self.debug:
            self.helper.save_debug("---DATE GETTER TOOL---")
            self.helper.save_debug(f'CURRENT DATE: {current_date}\n')

        self.state['context'] = [result]
        self.state['num_steps'] = num_steps

        return self.state

class InputTranslator(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are responsible of verifying if the USER_INPUT is in another language other
            than english, if so, translate the input. If the text is already in english
            simply output the same text. \n
            
            Your output must be a JSON object with two keys, 'language' and 'input', where
            'language' is the source language the user wrote and 'input' is the translated
            input. \n
        
            If the language is english, use ALWAYS the whole word to select it, never 'en',
            'eng', or any other variation. Use ALWAYS 'english'. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT : {user_input} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input"],
        )
    
    def execute(self) -> GraphState:
        user_input = self.state['user_input']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.ht_json_model | JsonOutputParser()

        llm_output = llm_chain.invoke({"user_input": user_input})
        translated_user_input = llm_output['input']
        source_language = llm_output['language']
        
        if self.debug:
            self.helper.save_debug("---TRANSLATE INPUT---")
            self.helper.save_debug(f'ORIGINAL INPUT: {user_input.rstrip()}')
            self.helper.save_debug(f'SOURCE LANGUAGE: {source_language}')
            if source_language.lower() != 'english':
                self.helper.save_debug(f'TRANSLATED INPUT:{translated_user_input.rstrip()}\n')
        
        self.state['num_steps'] = num_steps
        self.state['user_input'] = translated_user_input
        self.state['target_language'] = source_language
        
        return self.state

class ToolBypasser(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are part of the Energy System Insight Tool (ESIT), you are responsible for checking
            if the user trying to have a simple interaction with the tool instead of search for
            information or manipulate the model. \n

            Considering the USER_INPUT you should decide whether to route
            the user to the tool or simply bypass it to generate a simple answer to the user. \n
            
            The cases where you will bypass to the output are when USER_INPUT contains:
            - 'Hello!' or similars, without anything else;
            - 'How are you?' or similars, without anything else;
            - 'Who are you?' or similars;
            - 'What do you do?' or similars;
            - 'Thank you!' or similars;
            
            You must output a JSON with a single key 'is_conversation' containing exclusivelly the 
            selected type, which can be either true or false (use full lowercase for the boolean). \n
            
            Always use double quotes in the JSON object. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT : {user_input} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input"],
        )
    
    def execute(self) -> GraphState:
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        user_input = self.state['user_input']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        llm_output = llm_chain.invoke({"user_input": user_input})
        is_conversation = llm_output['is_conversation']
        if type(is_conversation) == str:
            is_conversation = True if is_conversation == 'true' else False
        elif type(is_conversation) != bool:
            is_conversation = False
        
        if self.debug:
            self.helper.save_debug("---TYPE IDENTIFIER---")
            self.helper.save_debug(f'BYPASS TO OUTPUT: {is_conversation}\n')
        
        self.state['is_conversation'] = is_conversation
        self.state['num_steps'] = num_steps
        
        return self.state

class TypeIdentifier(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are part of the Energy System Insight Tool (ESIT), your main task is to
            initiate the pipeline of the tool by deciding which type of request the user
            is trying to ask. You have three execution branches available, "energy_system",
            "mixed" and "general". \n

            "energy_system": the USER_INPUT is related to modeling. Unless the user specifies that the
            modeling he is talking about is not related to the model this tool is designed to
            analyze, you will use this branch. This branch allows you to complete diverse tasks
            related to Energy System modeling, like consulting information about the model,
            modifying values of the model, running it, comparing and plotting results, and also
            consults to the paper related to the model for more details. Anything on these lines
            should use this branch. You should also use this branch when the user talks about
            technical parameters, parametrization in general and asks about cientific data
            related to the development of the paper/model. Inputs without previous context pointing
            to the model that talk about changes, modifications, and other stuff related to 
            comparisons in different years normally are also related to the energy system. \n
            
            "mixed": the USER_INPUT would be identified as "energy_system", but there are references in
            the input that must be researched online to be able to gather the necessary context for the
            modelling tools to reach the user's goals. \n
            
            "general": the USER_INPUT is related to some generic topic, it may consist of one or more
            points that require searching for information. \n
            
            You must output a JSON with a single key 'input_type' containing exclusivelly the 
            selected type. \n
            
            You may also use the CHAT_HISTORY to get a better context of past messages exchanged
            between you and the user, and better understand what he wants in the current message. \n
            
            Always use double quotes in the JSON object. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT : {user_input} \n
            CHAT_HISTORY: {history} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input","history"],
        )
    
    def execute(self) -> GraphState:
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        user_input = self.state['user_input']
        history = self.state['history']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        llm_output = llm_chain.invoke({"user_input": user_input, "history": history})
        selected_type = llm_output['input_type']
        
        if self.debug:
            self.helper.save_debug("---TYPE IDENTIFIER---")
            self.helper.save_debug(f'USER INPUT: {user_input.rstrip()}')
            self.helper.save_debug(f'IDENTIFIED TYPE: {selected_type}\n')
            self.helper.save_debug(history)
        
        self.state['input_type'] = selected_type
        self.state['num_steps'] = num_steps
        
        return self.state
    
class InputConsolidator(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an expert at verifying if the USER_INPUT is complete on it's own or if it
            references data from the CHAT_HISTORY. \n
            
            You should output a JSON object with a single key 'consolidated_input', in which
            you have two possible situations.
            1. The USER_INPUT don't reference past messages and the following tools can use
            it as it is to execute their actions;
            2. The USER_INPUT references messages from CHAT_HISTORY, and this information is
            needed for the following tools. \n
            
            Your actions based on the situations are:
            1. Simply repeat the USER_INPUT as the 'consolidated_input';
            2. Build a new 'consolidated_input' adding the necessary information from the
            CHAT_HISTORY, while still keeping it as similar as possible with USER_INPUT. \n
            
            When modifying the input you must restrain yourself to substituting references to
            specific data by the data itself, for example, if the user talked about a specific
            parameter in the past message, and now said something like 'now modify it to X', you
            should substitute 'it' by the parameter name. Never add context related to results
            or general information. You should keep your rewriting as minimal as possible. \n
            
            Always use double quotes in the JSON object. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT : {user_input} \n
            CHAT_HISTORY: {history} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input", "history"],
        )
        
    def execute(self) -> GraphState:
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        user_input = self.state['user_input']
        history = self.state['history']
        num_steps = self.state['num_steps']
        num_steps += 1

        llm_output = llm_chain.invoke({"user_input": user_input, "history": history})
        consolidated_input = llm_output['consolidated_input']
        
        if self.debug:
            self.helper.save_debug("---INPUT CONSOLIDATOR---")
            self.helper.save_debug(f'USER INPUT: {user_input.rstrip()}')
            self.helper.save_debug(f'CONSOLIDATED INPUT: {consolidated_input}\n')
            
        self.state['consolidated_input'] = consolidated_input
        self.state['num_steps'] = num_steps
        
        return self.state

class ESNecessaryActionsSelector(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an expert at understanding the USER_INPUT and modifying ACTION_HISTORY to select
            the necessary actions to be taken to fulfill the user's request. \n
            
            ACTION_HISTORY provides you with the available actions to be taken, your job is to modify the
            dictionary to select the actions to take. You'll receive all actions as 'no', to select an action
            simply change that to 'yes', the rest of the pipeline will take care of executing them. \n
            
            The available actions are:
            'modify' - Whenever the model needs to be modified;
            'run' - Runs the model;
            'compare' - Compare the results of the runned model;
            'plot' - Plots the necessary data for user visualization;
            'consult' - Used to consult theoric information about the model.\n
            
            Here are the main possibilities that you have available from the user input:\n
            
            You will output ['modify': 'yes', 'run': 'yes', 'compare': 'yes', 'plot': 'yes', 'consult': 'no'] if USER_INPUT includes:
            - 'What if...';
            - 'What happens if...';
            - 'How the model reacts if...';
            - 'What changes if we...';
            - Other general cases where the USER_INPUT includes any type of 'What if?'\n

            You will output ['modify': 'no', 'run': 'no', 'compare': 'yes', 'plot': 'yes', 'consult': 'no'] if USER_INPUT includes:
            - 'What changed in...';
            - 'What happened to...';
            - 'Was there a variation in...';
            - Other general cases where the user asks about variation in results.\n
            
            You will output ['modify': 'no', 'run': 'no', 'compare': 'no', 'plot': 'yes', 'consult': 'no'] if USER_INPUT includes:
            - 'Show me...';
            - 'Plot the...';
            - 'Create the visualization of...';
            - Other general cases asking for data visualization.\n
            
            You will output ['modify': 'no', 'run': 'no', 'compare': 'no', 'plot': 'no', 'consult': 'yes'] if USER_INPUT includes:
            - 'How is ... modeled?';
            - 'Explain me about...';
            - 'How does the...';
            - 'What is the value of...';
            - Other cases where the user asks questions about the modelling and not the results.\n
            
            If USER_QUERY includes any of the action names, you should only output that actions as 'yes'.\n
            
            You must output a JSON with the modified dictionary. \n
            
            Always use double quotes in the JSON object. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT : {user_input} \n
            ACTION_HISTORY: {action_history} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input", "action_history"],
        )
        
    def execute(self) -> GraphState:
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        user_input = self.state['consolidated_input']
        action_history = self.state['action_history']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        llm_output = llm_chain.invoke({"user_input": user_input, "action_history": action_history})
        
        if self.debug:
            self.helper.save_debug("---ENERGY SYSTEM NECESSARY ACTIONS---")
            self.helper.save_debug(f'USER INPUT: {user_input.rstrip()}')
            self.helper.save_debug(f'NECESSARY ACTIONS: {llm_output}\n')
            
        self.state['action_history'] = llm_output
        self.state['num_steps'] = num_steps
        
        return self.state

class ESActionSelector(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return super().get_prompt_template()
        
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Processing')
        action_history = self.state['action_history']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        # These actions must be taken in this order of priority
        if action_history['modify'] in ['yes', 'repeat']:
            selected_action = 'modify'
        elif action_history['run'] == 'yes':
            selected_action = 'run'
        elif action_history['compare'] == 'yes':
            selected_action = 'compare'
        elif action_history['plot'] == 'yes':
            selected_action = 'plot'
        elif action_history['consult'] == 'yes':
            selected_action = 'consult'
        else:
            selected_action = 'no_action'
        
        if self.debug:
            self.helper.save_debug("---ENERGY SYSTEM ACTION SELECTOR---")
            self.helper.save_debug(f'SELECTED ACTION: {selected_action}\n')
            
        self.state['next_action'] = selected_action
        self.state['num_steps'] = num_steps
        
        return self.state

class QueryGenerator(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at developing a query to retrieve information from other means
            given the USER_INPUT and the CONTEXT with already known information. \n
            
            Using the CONTEXT to check which information is already available, you should analyze the
            USER_INPUT and decide what will be the next query for the available tools. You have 
            three options to gather information, online search, calculator and the aditional
            user's inputs. The identifiers of these options are ['web_search', 'calculator', 'user_input']. \n
            
            You'll also receive QUERY_HISTORY, use it to avoid repeating questions. If similar information was
            already searched at least 3 times without success, you may request for the user to input more
            information about what he wants. You may also check for the CHAT_HISTORY to verify what already
            happened in this session of the chat with the user, some information may be already present there. \n
            
            Restrain yourself as much as possible to request the user for more data, the only two cases you
            may do this are:
            1. The user asked for you to use information that only he knows (personal information for example);
            2. You already tried searching the information at least 3 times but couldn't find it, then you may
            prompt the user if the information could be provide more context for further research. If you couldn't
            find even having the necessary context, simply select 'no_action' as the selected tool. \n
            
            Also, never consider that you know how to to calculations, if there is any calculation in the 
            user request, build a query with it and pass it forward. For any mathematical problem you should use
            'calculator'. Be sure that you have all the necessary data before routing to this tool. \n
            
            You must output a JSON object with two keys, 'tool' and 'next_query'. 'tool' may contain one of the
            following ['web_search', 'calculator', 'user_input', 'no_action'] while 'next_query' is the next
            query to be processed by the tools. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            CONTEXT: {context} \n
            QUERY_HISTORY: {query_history} \n
            CHAT_HISTORY: {history} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input","context","query_history","history"],
        )
        
    def execute(self) -> GraphState:
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        ## Get the state
        user_input = self.state['user_input']
        context = self.state['context']
        query_history = self.state['query_history']
        history = self.state['history']
        num_steps = self.state['num_steps']
        num_steps += 1

        llm_output = llm_chain.invoke({"user_input": user_input,
                                   "context": context,
                                   "query_history": query_history,
                                   "history": history
                                   })
        
        if self.debug:
            self.helper.save_debug("---CONTEXT ANALYZER---")
            self.helper.save_debug(f'NEXT QUERY: {llm_output}\n')
            self.helper.save_debug(history)
        
        self.state['query_history'] = query_history + [llm_output['next_query']]
        self.state['next_query'] = llm_output
        self.state['num_steps'] = num_steps
        
        return self.state

class Mixed(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an expert at reading the user USER_INPUT and the available CONTEXT to decide if there
            is already enough information gathered to fulfill the energy system related command
            made by the user. \n
            
            You must be certain that you have all the data before deciding to send it to the
            modelling section of the pipeline. If any of the values asked by the user is not
            directly given by him, you can't consider the data complete unless you have the
            desired value in the CONTEXT. \n
            
            You may also check for the CHAT_HISTORY to verify what already happened in this session
            of the chat with the user. \n

            You must output a JSON object with a single key 'is_data_complete' containing a boolean
            on whether you have enough data for the user's request or not. \n
            
            Always use double quotes in the JSON object. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT : {user_input} \n
            CONTEXT: {context} \n
            CHAT_HISTORY: {history} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input","context","history"],
        )
        
    def execute(self) -> GraphState:
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        user_input = self.state['user_input']
        context = self.state['context']
        history = self.state['history']
        num_steps = self.state['num_steps']
        num_steps += 1

        llm_output = llm_chain.invoke({"user_input": user_input, "context": context, "history": history})
        is_data_complete = llm_output['is_data_complete']
        
        if self.debug:
            self.helper.save_debug("---TOOL SELECTION---")
            self.helper.save_debug(f'USER INPUT: {user_input.rstrip()}')
            self.helper.save_debug(f'CONTEXT: {context}')
            self.helper.save_debug(f'DATA IS COMPLETE: {is_data_complete}\n')
            
        self.state['is_data_complete'] = is_data_complete
        self.state['num_steps'] = num_steps
        
        return self.state

class ContextAnalyzer(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an expert at analyzing the available CONTEXT and CHAT_HISTORY to decide if the available
            information is already enough to answer the question asked in USER_INPUT. \n

            It's important to know that you are the context analyzer for a branch of a tool, and this branch is
            responsible to gather general information to be used in the modeling context. If the USER_INPUT
            is related to modeling and uses information available at the data sources to apply modeling actions
            and all this information is available, then you should consider that the context is ready. Other
            tools will be responsible of using this data for modeling. \n
            
            If there is nothing related to modeling you can simply define it as ready when all information is
            gathered. \n
            
            Your output is a JSON object with a single key 'ready_to_answer', where you can use either true
            or false (always write it in lowercase). \n
            
            Always use double quotes in the JSON object. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            CONTEXT : {context} \n
            CHAT_HISTORY: {history} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input","context","history"]
        )
        
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Processing')
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        user_input = self.state['consolidated_input']
        context = self.state['context']
        history = self.state['history']
        num_steps = self.state['num_steps']
        num_steps += 1

        llm_output = llm_chain.invoke({"user_input": user_input, "context": context, "history": history})
        
        if self.debug:
            self.helper.save_debug("---TOOL SELECTION---")
            self.helper.save_debug(f'READY TO ANSWER: {llm_output["ready_to_answer"]}\n')
        
        self.state['next_query'] = llm_output
        self.state['num_steps'] = num_steps
        
        return self.state

class ResearchInfoWeb(ResearchAgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a master at working out the best keywords to search for in a web search to get the best info for the user.

            Given the QUERY, work out the best search queries that will find the info requested by the user
            The queries must contain no more than a phrase, since it will be used for online search. 

            Return a JSON with a single key 'keywords' with at most 3 different search queries.
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            QUERY: {query} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["query"],
        )
        
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Searching info in the internet')
        if self.debug:
            self.helper.save_debug("---RESEARCH INFO SEARCHING---")
            
        query = self.state['next_query']['next_query']
        context = self.state['context']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        prompt = self.get_prompt_template()
        answer_analyzer_prompt = self.get_answer_analyzer_prompt_template()
        
        llm_chain = prompt | self.json_model | JsonOutputParser()
        answer_analyzer_chain = answer_analyzer_prompt | self.chat_model | StrOutputParser()

        # Web search
        keywords = llm_chain.invoke({"query": query, "context": context})
        keywords = keywords['keywords']
        full_searches = []
        for idx, keyword in enumerate(keywords):
            temp_docs = self.web_tool.execute(keyword)
            if type(temp_docs) == list:
                for d in temp_docs:
                    web_results = f'Source: {d["url"]}\n{d["content"]}\n'
                web_results = Document(page_content=web_results)
            elif type(temp_docs) == dict:
                web_results = f'\nSource: {d["url"]}\n{d["content"]}'
                web_results = Document(page_content=web_results)
            else:
                web_results = 'No results'
            if self.debug:
                self.helper.save_debug(f'KEYWORD {idx}: {keyword}')
                self.helper.save_debug(f'RESULTS FOR KEYWORD {idx}: {web_results}')
            if full_searches is not None:
                full_searches.append(web_results)
            else:
                full_searches = [web_results]

        processed_searches = answer_analyzer_chain.invoke({"query": query, "search_results": full_searches, "context": context})
        
        if self.debug:
            self.helper.save_debug(f'FULL RESULTS: {full_searches}\n')
            self.helper.save_debug(f'PROCESSED RESULT: {processed_searches}\n')
        
        self.state['context'] = context + [processed_searches]
        self.state['num_steps'] = num_steps
        
        return self.state

class Calculator(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an specialist at identifying the correct operation and operands to build
            a JSON object with the correct equation to be solved, you will get the information of
            what equation to generate from the given QUERY. \n
            
            You must output a JSON object with a single key 'equation', that must be in a format
            solvable with Python, since we'll solve this using the eval() function. You don't need
            to import anything or prepare the environment in any way, the environment is ready,
            you just need to build the equation in a format recognizable by python, using python
            mathematical functions if necessary. \n
            
            NEVER USE IMPORT STATEMENTS, YOU MUST ASSUME THAT ALL NECESSARY LIBRARIES ARE ALREADY
            IMPORTED, JUST USE THEM. \n
            
            If you need information that you don't know to build the equation you may find that
            in the CONTEXT, so check it for more details on the data. \n
            
            If you need intermediate results, you can calculate only them and leave the final
            result for later, you're part of a iterative tool that may be called as many times
            as needed to find the final result. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            QUERY: {query} \n
            CONTEXT: {context}
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["query","context"],
        )
        
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Calculating result')
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
    
        query = self.state['next_query']['next_query']
        context = self.state['context']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        llm_output = llm_chain.invoke({"query": query, "context": context})
        equation = llm_output['equation']
        
        if self.debug:
            self.helper.save_debug("---CALCULATOR TOOL---")
            self.helper.save_debug(f'EQUATION: {equation}')
        
        result = eval(equation)
        
        str_result = f'{equation} = {result}'
        
        if self.debug:
            self.helper.save_debug(f'RESULT: {str_result}\n')
            
        self.state['context'] = context + [str_result]
        self.state['num_steps'] = num_steps
        
        return self.state

class RunModel(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return super().get_prompt_template()
    
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Running model')
        action_history = self.state['action_history']
        context = self.state['context']
        model_modified = self.state['model_modified']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        if self.debug:
            self.helper.save_debug('---SIMULATION RUNNER---')
        
        if model_modified:
            try:
                model_name = self.mod_model.split('/')[-1]
                model_name = model_name[:-5]
                scenario = 'Base'
                techmap_dir_path = Path("./CESM").joinpath('Data', 'Techmap')
                ts_dir_path = Path("./CESM").joinpath('Data', 'TimeSeries')
                runs_dir_path = Path("./CESM").joinpath('Runs')
                
                # Create a directory for the model if it does not exist
                db_dir_path = runs_dir_path.joinpath(model_name+'-'+scenario)
                if not runs_dir_path.exists():
                    os.mkdir(runs_dir_path)
                
                if not os.path.exists(db_dir_path):
                    os.mkdir(db_dir_path)

                # Create and Run the model
                conn = sqlite3.connect(":memory:")
                parser = Parser(model_name, techmap_dir_path=techmap_dir_path, ts_dir_path=ts_dir_path, db_conn = conn, scenario = scenario)

                # Parse
                if self.debug:
                    self.helper.save_debug("\n#-- Parsing started --#")
                parser.parse()

                # Build
                if self.debug:
                    self.helper.save_debug("\n#-- Building model started --#")
                model_instance = Model(conn=conn)

                # Solve
                try:
                    if self.debug:
                        self.helper.save_debug("\n#-- Solving model started --#")
                    
                    with open(self.helper.get_debug_log_path(), 'a') as f:
                        sys.stdout = f
                        model_instance.solve()
                    sys.stdout = sys.__stdout__
                except Exception as e:
                    sys.stdout = sys.__stdout__
                    self.helper.save_debug(e)

                # Save
                if self.debug:
                    self.helper.save_debug("\n#-- Saving model started --#")
                model_instance.save_output()

                
                db_path = db_dir_path.joinpath('db.sqlite')
                if db_path.exists():
                    # Delete the file using unlink()
                    db_path.unlink()
                
                # write the in-memory db to disk
                disk_db_conn = sqlite3.connect(db_path)
                conn.backup(disk_db_conn)
                sim_status = 'runned'
                self.helper.save_simulation_status('runned')
                result = 'The model has runned successfully!'
            except Exception as e:
                # TODO set a better way of defining infeasibility
                self.helper.save_debug(e)
                sim_status = 'infeasible'
                result = 'Something went wrong and the model has not completely runned.'
        else:
            sim_status = 'no_run'
            result = 'There were no modifications to the model, no new simulation required'
            
        if self.debug:
            self.helper.save_debug(f'SIMULATION RESULTS: {result}\n')
        
        action_history['run'] = 'done'
        self.state['action_history'] = action_history
        self.state['context'] = context + [result]
        self.state['sim_status'] = sim_status
        self.state['num_steps'] = num_steps
        
        return self.state

class ModifyModel(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return super().get_prompt_template()
    
    def get_sheet_selector_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at identifying the correct sheet to be modified based on the user's USER_INPUT. \n
            
            You must output a JSON object with a single key 'sheet', that should receive 'conversionsubprocess'.
            You should only output 'scenario' instead if the user explicitly talks about co2 limit, co2 price
            or discount rate, these terms must be included in the query. \n
            
            Remember that for the sheet to be 'scenario' you need to receive explicitly the defined terms from
            the USER_INPUT. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input"],
        )
    
    def get_scenario_param_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at modifying the scenario of a model based on the user's USER_INPUT. \n
            
            There are three parameters you can change: discount_rate, annual_co2_limit and co2_price.
            In SCEN_PARAMS you have a dictionary with the current values for each of them, and you should
            decide how to change them to fulfill the user's request. \n
            
            YOU SHOULD ONLY MODIFY THE PARAMETERS IF THE USER CLEARLY STATES THAT HE WANTS ONE OF THE
            SCENARIO PARAMETERS TO BE MODIFIED. \n
            
            You must output a JSON object with a single element 'new_values' that contains the modified
            values for the three in a three elements list. If you haven't modified some of them, simply
            output the same value you received. The order should be the same as the input.\n
            
            For the anual co2 limit, if the user request a change after a specific year you should always
            add the next year to the yearly list as a starting point for the modification, and keep the remaining
            years, just changing their values. Also, for this list the only possibility are numbers for the 
            values, not 'infinity' or anything like that. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            SCEN_PARAMS: {scen_params}
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input", "scen_params"],
        )
    
    def get_params_general_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at identifying the correct way to modify a model based on the user's USER_INPUT. \n
            
            You are part of a tool where there are two other agents ready to execute their specific tasks
            related to model modification. Your goal is to identify which one to use depending on the user's
            input. \n
            
            The two possibilities are:
            1. The user provides one or more specific instructions of modifications that should be done to the model,
            with the selection of specific values for each modification;
            2. The user asks for a specific scenario where the tool should decide which parameters and subprocesses
            should be modified, as well as deciding the correct value to change the combinations to. \n
            
            Your output must be a JSON object with a single key called 'parametrization_type', where your options are
            'defined' for the first case and 'undefined' for the second. These are your only possibilities. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input"],
        )

    def get_select_cs_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at analizing the USER_INPUT and determining the correct conversion
            suprocesses that should be modified in a energy system to achieve the user request. \n
            
            Your selection should come from the CONVERSION_SUBPROCESSES list provided, you are not
            allowed to guess names of conversion subprocesses or modify the entries of the list before
            returning them. You MUST use them as they are in the provided list. \n
            
            There are two possibilities of ways for you to choose the correct subprocesses:
            1. The user specified what he wants to modify in the model, in this case you should try
            to match what was requested and nothing else;
            2. The user requested a general scenario in which you will be responsible to decide
            which subprocesses should be modified. \n
            
            Your output must be a JSON object with a single key 'cs_selection' that will contain a
            single list with all entries that you selected from CONVERSION_SUBPROCESSES. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            CONVERSION_SUBPROCESSES: {cs_list} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input","cs_list"],
        )
        
    def get_select_params_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at analizing the USER_INPUT and determining the correct parameters
            that should be modified in a energy system to achieve the user request. \n
            
            Your selection should come from the PARAMETERS dictionary provided, you are not
            allowed to guess names of parameters or modify the entries of the list before
            returning them. You MUST use them as they are in the provided list. \n
            
            The PARAMETERS dictionary shows for each conversion subprocess the available parameters,
            you should select from among the ones available for that specific subprocess. The format
            in which each parameter is shown is 'param_name - description', you must output only the
            param_name, never the description. \n
            
            There are two possibilities of ways for you to choose the correct parameters:
            1. The user specified what he wants to modify in the model, in this case you should try
            to match what was requested and nothing else;
            2. The user requested a general scenario in which you will be responsible to decide
            which parameters should be modified. \n
            
            Your output must be a JSON object with a single key 'param_selection' that will contain a
            single dictionary where the keys are each of the conversion subprocesses and the values
            are the list of selected parameters for each of them. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            PARAMETERS: {param_list} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input","param_list"],
        )

    def get_new_values_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at defining new values for model parameters based on the USER_INPUT. \n
            
            You have access to all values to be modified in CURRENT_VALUES, they are organized as a
            dictionary, where each key represents a conversion subprocess and the value of each key
            is a list of lists containing all of the current values. Each sublist of this list is
            organized as follows: [param_name, current_value, unit]. \n
            
            Your job is to understand the USER_INPUT to modify the 'current_value' in this structure.
            You will output a JSON object that contains two keys, 'success' and 'values', where 'success'
            should receive a boolean (always lower case) and 'values' should receive the same structure
            as CURRENT_VALUES but with 'current_value' updated for the new value in all entries. \n
            
            You have some options of how to update the value based on USER_INPUT:
            1. The user specified a value, in this case you need to use the exact values that were asked
            by the user;
            2. The user didn't specify a value but gave some kind of idea of the modification, for example
            asked you to double the value, to reduce by 30%, to remove (in this case the value should go to 0),
            in this case you need to figure out by how much the value will be modified;
            3. The user didn't specify a value but asked for some indirect value related to other piece of
            information, in this case you can check CONTEXT for the necessary information;
            4. In case you couldn't categorize the request as any of the past 3 cases, you should mark
            'success' as false and 'values' as an empty dict. \n
            
            There are two types of values that you can receive, numbers or lists. If the value is a number
            it's simple, the number represents the value of that parameter, however, if you receive a list
            (which will be in form of a string and should also be outputed as a string) it represents the
            variation of the parameter by year, for example '[2015 10; 2030 20; 2050 40]'. You may modify
            the available years and the values, as long as you keep the format and the output as string. \n
            
            If you need to output any value as 'nan', 'null', empty or similars, you should instead use
            and empty string as the output for that value. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            CONTEXT: {context} \n
            CURRENT_VALUES: {current_values} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input","context","current_values"],
        )

    def get_mod_ready_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at identifying if there is any more modifications to be done to a given model. \n
            
            You will receive the USER_INPUT that trigered the modification process and the MODIFICATIONS that were
            already done to the model. Your goal is determine if the necessary modifications were already done
            or if there are still parameters to be modified in the model. \n
            
            Your output must be a JSON object with a single key 'ready' with a boolean (always write booleans
            with lowercase) that states whether the modifications are ready or if the model needs to be further
            modified. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            MODIFICATIONS: {modifications}
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input", "modifications"],
        )

    def execute(self) -> GraphState:
        self.helper.save_chat_status('Modifying model')
        if self.debug:
            self.helper.save_debug('---MODEL MODIFIER---')
        
        helper = HelperFunctions()
        
        # Define prompts
        sheet_selector_prompt = self.get_sheet_selector_prompt_template()
        scenario_param_prompt = self.get_scenario_param_prompt_template()
        mod_ready_prompt = self.get_mod_ready_prompt_template()
        params_general_prompt = self.get_params_general_prompt_template()
        select_cs_prompt = self.get_select_cs_prompt_template()
        select_params_prompt = self.get_select_params_prompt_template()
        new_values_prompt = self.get_new_values_prompt_template()

        # Define chains
        sheet_selector_chain = sheet_selector_prompt | self.json_model | JsonOutputParser()
        scenario_param_chain = scenario_param_prompt | self.json_model | JsonOutputParser()
        mod_ready_chain = mod_ready_prompt | self.json_model | JsonOutputParser()
        params_general_chain = params_general_prompt | self.json_model | JsonOutputParser()
        select_cs_chain = select_cs_prompt | self.json_model |JsonOutputParser()
        select_params_chain = select_params_prompt | self.json_model |JsonOutputParser()
        new_values_chain = new_values_prompt | self.json_model | JsonOutputParser()
        
        # Define inputs from state
        user_input = self.state['consolidated_input']
        context = self.state['context']
        action_history = self.state['action_history']
        num_steps = self.state['num_steps']
        num_steps += 1

        # Get list of params, conversion subprocesses and dict with the scenario params
        params, CSs = helper.get_params_and_cs_list(self.base_model)
        scen_params = helper.get_scenario_params(self.base_model)
        
        # If the modification already happened, go directly to conversionsubprocess
        if action_history['modify'] == 'repeat':
            sheet = 'conversionsubprocess'
        else:
            # Selects the sheet that must be modified (scenario or cs)
            sheet_selection = sheet_selector_chain.invoke({"user_input": user_input})
            sheet = sheet_selection['sheet']
        
        # Load workbook for modifications
        workbook = load_workbook(filename=self.base_model)
        
        if sheet == 'scenario':
            # Modify the scenario parameters
            new_params = scenario_param_chain.invoke({"user_input": user_input,
                                                      "scen_params": scen_params})
            
            workbook, result = helper.modify_scenario_sheet(workbook, new_params)
            if len(result) == 0:
                message = 'Nothing was modified'
            else:
                message = f'The model was modified as follows:{result}'
            
            modifications_ready = mod_ready_chain.invoke({"user_input": user_input,
                                                          "modifications": message})
            if modifications_ready['ready']:
                action_history['modify'] = 'done'
            else:
                action_history['modify'] = 'repeat'
        else:
            try:
                # Defines the type of parametrization (defined or undefined)
                parametrization_type = params_general_chain.invoke({"user_input": user_input})
                parametrization_type = parametrization_type['parametrization_type']

                # Selects the conversion subprocesses that match the user input
                cs_selection = select_cs_chain.invoke({"user_input": user_input, "cs_list": CSs})
                cs_selection = cs_selection['cs_selection']

                if parametrization_type == 'defined':
                    # Set the available parameters as being all of them for each conversion subprocess
                    available_parameters = {cs: params for cs in cs_selection}
                else:
                    # Set the available parameters as being only the ones that are already filled for each cs
                    available_parameters = helper.get_populated_params_and_cs_list(self.base_model, cs_selection)

                # Selects the matching parameter from the set of available ones
                param_selection = select_params_chain.invoke({"user_input": user_input, "param_list": available_parameters})
                param_selection = param_selection['param_selection']

                # Gets the current values for the selected pairs of cs and parameter
                current_values = helper.get_values(self.base_model, param_selection)

                # Feed the current data to the model, it will output the updated values in the same format
                new_params = new_values_chain.invoke({"user_input": user_input, "context": context, "current_values": current_values})
                
                # IF the agent indicates that the modification was successfull, then apply them to the sheet
                if new_params['success']:
                    workbook, result = helper.modify_cs_sheet(workbook, new_params)
                    if len(result) == 0:
                        message = 'Nothing was modified'
                    else:
                        message = f'The model was modified as follows:{result}'
                else:
                    message = ['Failed to generate the correct modified set of parameters.']
            except Exception as e:
                self.helper.save_debug(e)
                message = ['Failed to modify the model, probably nothing to change']
            
            action_history['modify'] = 'done'
                        
        try:
            # Try to save, if it saves and the scenario was modified, mark it as modified
            workbook.save(filename=self.mod_model)
            self.state['model_modified'] = True
        except Exception as e:
            self.helper.save_debug(e)
            message = ['Failed to save the modifications']
            action_history['modify'] = 'done'
        
        if self.debug:
            self.helper.save_debug(f'SELECTED PARAMS + VALUES: {new_params}')
            self.helper.save_debug(f'FINAL RESULTS OF MODIFICATION:\n{message}\n')
        
        self.state['num_steps'] = num_steps
        self.state['context'] = context + [message]
        self.state['action_history'] = action_history
        
        return self.state

class InfoTypeIdentifier(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are part of a modelling insights system, in this system we have two sources of
            information, the paper that explains the model that we use and the model data itself.
            You are an expert at defining from which of these two the information asked by the
            user should be gathered. \n
            
            You have access to the USER_INPUT, and based on that you should decide if the output
            is 'model' or 'paper'. \n
            
            Normally, the output will be 'model' if the user asks for the value of a specific
            parameter, or if the user asks about components that are modeled (such as commodity,
            conversion processes, conversion subprocesses and scenario information) in a general
            way (not asking for detailed modelling information). \n
            
            On other cases you should use 'paper', that's where the theoric information behind
            the model lies, information such as intrinsec details of the modeling process,
            constraints, parameter relationships, details of the types of parameters and 
            implementation details. \n

            Return a JSON with a single key 'type' that contains either 'model' or 'paper'. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input"],
        )
    
    def execute(self) -> GraphState:
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.json_model | JsonOutputParser()
        
        user_input = self.state['consolidated_input']
        num_steps = self.state['num_steps']
        num_steps += 1

        llm_output = llm_chain.invoke({"user_input": user_input})
        
        if self.debug:
            self.helper.save_debug("---INFO TYPE IDENTIFIER---")
            self.helper.save_debug(f'RETRIEVAL TYPE: {llm_output["type"]}\n')
        
        self.state['retrieval_type'] = llm_output['type']
        self.state['num_steps'] = num_steps
        
        return self.state

# TODO check the answer analyzer prompt
# TODO create chain to decide whether to search information on the paper or on the CESM documentation

class ResearchInfoRAG(ResearchAgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a master at working out the best questions to ask our knowledge agent to get 
            the best info for the customer. \n

            Given the QUERY, work out the best questions that will find the best info for 
            helping to write the final answer. Write the questions to our knowledge system not 
            to the customer. \n

            Return a JSON with a single key 'questions' with no more than 3 strings of and no 
            preamble or explaination. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            QUERY: {query} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["query"],
        )
        
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Consulting the source paper')
        question_rag_prompt = self.get_prompt_template()
        answer_analyzer_prompt = self.get_answer_analyzer_prompt_template()
        
        question_rag_chain = question_rag_prompt | self.json_model | JsonOutputParser()
        answer_analyzer_chain = answer_analyzer_prompt | self.chat_model | StrOutputParser()
        
        if self.debug:
            self.helper.save_debug("---RAG PDF PAPER RETRIEVER---")
            
        query = self.state['consolidated_input']
        context = self.state['context']
        action_history = self.state['action_history']
        num_steps = self.state['num_steps']
        num_steps += 1

        questions = question_rag_chain.invoke({"query": query})
        questions = questions['questions']

        rag_results = []
        for idx, question in enumerate(questions):
            temp_docs = self.retriever.execute(question)
            if self.debug:
                self.helper.save_debug(f'QUESTION {idx}: {question}')
                self.helper.save_debug(f'ANSWER FOR QUESTION {idx}: {temp_docs.response}')
            question_results = question + '\n\n' + temp_docs.response + "\n\n\n"
            if rag_results is not None:
                rag_results.append(question_results)
            else:
                rag_results = [question_results]
        if self.debug:
            self.helper.save_debug(f'FULL ANSWERS: {rag_results}\n')
        
        processed_searches = answer_analyzer_chain.invoke({"query": query, "search_results": rag_results, "context": context})
        # TODO find a way of referencing the used pdfs
        result = f'Source: PDF paper \n{query}: \n{processed_searches}'
        
        action_history['consult'] = 'done'
        self.state['action_history'] = action_history
        self.state['context'] = context + [result]
        self.state['num_steps'] = num_steps
        
        return self.state
    
# TODO re-think the way the info flows from and to here

class ConsultModel(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are a specialist at identifying from the user's USER_INPUT the correct consultation that
            should be made to the model to return the required information to the user. \n
            
            You'll receive CONVERSION_PROCESSES, CONVERSION_SUBPROCESSES, PARAMETERS and SCENARIO_INFO
            as context. \n
            
            CONVERSION_PROCESSES is a list with the name of all conversion processes. They are the general 
            processes of energy conversion and generation, better fragmentalized in CONVERSION_SUBPROCESSES. \n
            
            CONVERSION_SUBPROCESSES is presented to you in the format "'cp'@'cin'@'cout' - 'description'" and
            describes the processes of energy exchange, such as gas to heat, or gas to electricity, for each
            of the conversion processes. \n
            
            PARAMETERS is a list with all parameters that can be modeled for the conversion subprocesses along
            with their descriptions. \n
            
            SCENARIO_INFO is a dictionary with the relevant information about the scenario to consider
            if asked by the user. It contains the discount rate of the model, the anual limits of CO2 emission
            and the defined price for CO2. \n
            
            Beyond the already available information, you can ask for two other types of information,
            ['yearly_variation', 'cs_param_selection']. \n
            
            - yearly_variation: provides you with all parameters of specific conversion subprocesses that
            have yearly variation. The parameter is displayed as 'conversion_process'@'cin'@'cout'@'parameter', and
            the value comes as a list with years and values, such as [2016 10; 2020 20; 2030 30] for example. This
            would show you yearly variation of parameters for different conversion subprocesses;
            - cs_param_selection: allows you to choose a set of conversion subprocesses and parameters
            to get their values from the model. \n
            
            Your only output should be a JSON with three keys, 'consult_type', 'cs' and 'param'. 'consult_type' must
            be 'yearly_variation' or 'cs_param_selection, 'cs' and 'param' are optional, and should only be used
            in the 'cs_param_selection' case. \n
            
            Whenever you need to use 'cs_param_selection', you should get 'cs' from CONVERSION_SUBPROCESS and
            'param' from PARAMETERS. You MUST NOT modify the entries, you should use them exactly as they are given
            to you in the input lists, you are also NOT ALLOWED to guess CSs or params, use only the ones
            available in the lists. The values in 'cs' and 'param' should always be lists, even if you want
            to know a single value. \n
            
            Always use double quotes in the JSON object. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            CONVERSION_PROCESSES: {cp} \n
            CONVERSION_SUBPROCESSES: {cs} \n
            PARAMETERS: {param} \n
            SCENARIO_INFO: {scen_infos} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input","cp","cs","param","scen_infos"],
        )
    
    def get_summarizing_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
                You are an specialist at sumarizing model information based on the user's USER_INPUT. \n
                
                You'll receive CONVERSION_PROCESSES, CONVERSION_SUBPROCESSES, PARAMETERS SCENARIO_INFO
                and CONSULTED_INFO as context. \n
                
                CONVERSION_PROCESSES is a list with the name of all conversion processes. They are the general 
                processes of energy conversion and generation, better fragmentalized in CONVERSION_SUBPROCESSES. \n
                
                CONVERSION_SUBPROCESSES is presented to you in the format "'cp'@'cin'@'cout' - 'description'" and
                describes the processes of energy exchange, such as gas to heat, or gas to electricity, for each
                of the conversion processes. \n
                
                PARAMETERS is a list with all parameters that can be modeled for the conversion subprocesses along
                with their descriptions. \n
                
                SCENARIO_INFO is a dictionary with the relevant information about the scenario to consider
                if asked by the user. It contains the discount rate of the model, the anual limits of CO2 emission
                and the defined price for CO2. \n
                
                CONSULTED_INFO contains possible extra information about the model that you may need to evaluate
                the answer. \n
                
                Your goal with all this context is to sumarize the information requested by the user about the
                model based on the USER_INPUT. Make the summarization the most natural possible, without citing the name
                of the source inputs. \n

                <|eot_id|><|start_header_id|>user<|end_header_id|>
                USER_INPUT: {user_input} \n
                CONVERSION_PROCESSES: {cp} \n
                CONVERSION_SUBPROCESSES: {cs} \n
                PARAMETERS: {param} \n
                SCENARIO_INFO: {scen_infos} \n
                CONSULTED_INFO: {info} \n
                Answer:
                <|eot_id|>
                <|start_header_id|>assistant<|end_header_id|>
                """,
                input_variables=["user_input","cp","cs","param","scen_infos","info"],
        )
    
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Consulting the model data')
        user_input = self.state['consolidated_input']
        context = self.state['context']
        action_history = self.state['action_history']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        consult_prompt = self.get_prompt_template()
        sum_up_prompt = self.get_summarizing_prompt_template()
        
        consult_model_chain = consult_prompt | self.json_model | JsonOutputParser()
        sum_up_info_chain = sum_up_prompt | self.chat_model | StrOutputParser()
        
        helper = HelperFunctions()

        params, CSs = helper.get_params_and_cs_list(self.base_model)
        CPs = helper.get_conversion_processes(self.base_model)
        scen_infos = helper.get_scenario_params(self.base_model)
        
        llm_output = consult_model_chain.invoke({"user_input": user_input,
                                             "cp": CPs,
                                             "cs": CSs,
                                             "param": params,
                                             "scen_infos": scen_infos})
        
        if self.debug:
            self.helper.save_debug('---CONSULT MODEL---')
            self.helper.save_debug(f'ANALYZER OUTPUT: {llm_output}')

        info = helper.consult_info(llm_output, self.base_model)
        
        summed_up_info = sum_up_info_chain.invoke({"user_input": user_input,
                                                   "cp": CPs,
                                                   "cs": CSs,
                                                   "param": params,
                                                   "scen_infos": scen_infos,
                                                   "info": info})
        
        if self.debug:
            self.helper.save_debug(f'GATHERED INFO ABOUT THE MODEL: {summed_up_info}')
        
        self.state['num_steps'] = num_steps
        action_history['consult'] = 'done'
        self.state['action_history'] = action_history
        self.state['context'] = context + [summed_up_info]
        
        return self.state

class CompareModel(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return super().get_prompt_template()
    
    def get_analysis_type_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
                You are an specialist at checking the USER_INPUT to decide if the user wants to know the comparison
                between the original model and the modified model, which can be mainly identified by the user asking
                for some parameters to be modified in the model, or if the user wants to check result variation
                within different years of a single model, which normally happens when the user just ask about differences
                in the results compared to now or to a specific year. \n
                
                Your only output should be a JSON object with a single key 'type', that can be either 'model_diff'
                or 'yearly_diff'. \n
                
                Always use double quotes in the JSON object. \n

                <|eot_id|><|start_header_id|>user<|end_header_id|>
                USER_INPUT: {user_input} \n
                Answer:
                <|eot_id|>
                <|start_header_id|>assistant<|end_header_id|>
                """,
                input_variables=["user_input"],
            )

    def get_results_diff_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
                You are an specialist at analyzing the variations in the model's output variables to summarize
                the most relevant information given the USER_INPUT and the OUTPUT_VARIATIONS. \n
                
                OUTPUT_VARIATIONS is a dictionary that contains the percentual variation of the output variables
                of a model after the user modified it and ran it again. The first layer has the output variables
                as the key, the second has each year of the simulation as a key, and finally, the values of these
                are the the variations for each subprocess (represented by a combination of 
                'conversion_process'@'commodity_in'@'commodity_out'). \n
                
                You also have access to the capital cost (CAPEX), operational cost (OPEX) and the total cost
                (TOTEX = CAPEX + OPEX) of the model in COSTS_VARIATION. However, nevem mention them in the
                output if the variation was -100%, something probably went wrong and that value is not valid. \n
                
                You must consider that the model was modified to account for the user's modification request shown
                in USER_INPUT, then the modified model was simulated and the variation you have available in
                OUTPUT_VARIATIONS is the variation between the original model and the model with the modifications
                requested by the user. Using this context you should be able to give the user the main insights
                about the scenario that he requested. \n

                <|eot_id|><|start_header_id|>user<|end_header_id|>
                USER_INPUT: {user_input} \n
                OUTPUT_VARIATIONS: {output_variations} \n
                COSTS_VARIATION: {costs_variation} \n
                Answer:
                <|eot_id|>
                <|start_header_id|>assistant<|end_header_id|>
                """,
                input_variables=["user_input","output_variations","costs_variation"],
            )

    def get_yearly_diff_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
                You are an specialist at analyzing the variations in the model's output variables to summarize
                the most relevant information given the USER_INPUT and the YEARLY_VARIATIONS. \n
                
                YEARLY_VARIATIONS is a dictionary that contains the percentual variation of the output variables
                throughout the different years of the model results. The first layer has the output variables
                as the key, the second has each conversion subprocess (represented by a combination of 
                'conversion_process'@'commodity_in'@'commodity_out') of the simulation as a key, and finally,
                the values of these are the the variations for each year. \n
                
                Using this context you should be able to give the user the main insights about the scenario
                yearly results comparison that he requested. \n

                <|eot_id|><|start_header_id|>user<|end_header_id|>
                USER_INPUT: {user_input} \n
                YEARLY_VARIATIONS: {yearly_variations} \n
                Answer:
                <|eot_id|>
                <|start_header_id|>assistant<|end_header_id|>
                """,
                input_variables=["user_input","yearly_variations"],
            )

    def execute(self) -> GraphState:
        self.helper.save_chat_status('Comparing results')
        user_input = self.state['consolidated_input']
        context = self.state['context']
        action_history = self.state['action_history']
        sim_status = self.state['sim_status']
        num_steps = self.state['num_steps']
        num_steps += 1

        runs_dir_path = 'CESM/Runs'
        
        simulation_base = f'{self.base_model.split("/")[-1][:-5]}'
        simulation_new = f'{self.mod_model.split("/")[-1][:-5]}'
        # simulation_new = 'DEModelMock' # For debugging

        # Set interest variables
        non_t_variables = ["Cap_new",
                           "Cap_active",
                           "Cap_res",
                           "Eouttot",
                           "Eintot",
                           "E_storage_level_max"]

        cost_vars = ["OPEX",
                     "CAPEX",
                     "TOTEX"]

        df_base = self.helper.load_results(runs_dir_path, simulation_base, non_t_variables)
        df_new = self.helper.load_results(runs_dir_path, simulation_new, non_t_variables)
        df_costs_base = self.helper.load_results(runs_dir_path, simulation_base, cost_vars)
        df_costs_new = self.helper.load_results(runs_dir_path, simulation_new, cost_vars)

        # Fill missing years with zeroes so we can compare the dataframes correctly
        # Doing that because the results can have different sets of CSs and years
        df_base, df_new = self.helper.fill_empty_rows(df_base, df_new)
        
        if sim_status == 'runned':
            prompt = self.get_analysis_type_prompt_template()
            llm_chain = prompt | self.json_model | JsonOutputParser()
            
            llm_output = llm_chain.invoke({"user_input": user_input})
            type = llm_output['type']
        else:
            type = 'yearly_diff'

        # TODO implement verification of status beyond runned ('infeasible', 'failed', 'no_run')
        if type == 'model_diff':
            variations_costs = (df_costs_new / df_costs_base - 1) * 100
            variations_dict = self.helper.get_models_variation(df_base, df_new, non_t_variables)

            prompt = self.get_results_diff_prompt_template()
            llm_chain = prompt | self.ht_model | StrOutputParser()

            llm_output = llm_chain.invoke({"user_input": user_input, "output_variations": variations_dict, "costs_variation": variations_costs})
        else:
            variations_dict = self.helper.get_yearly_variations_from_results(df_base, non_t_variables)
            
            prompt = self.get_yearly_diff_prompt_template()
            llm_chain = prompt | self.ht_model | StrOutputParser()

            llm_output = llm_chain.invoke({"user_input": user_input, "yearly_variations": variations_dict})

        if self.debug:
            self.helper.save_debug("---COMPARE RESULTS---")
            self.helper.save_debug(f'ANALYSIS TYPE: {type}')
            self.helper.save_debug(f'ANALYSIS: {llm_output}\n')
        
        self.state['context'] = context + [llm_output]
        action_history['compare'] = 'done'
        self.state['action_history'] = action_history
        self.state['num_steps'] = num_steps
        
        return self.state

class PlotModel(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return super().get_prompt_template()
    
    def get_select_plot_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an specialist at deciding the correct type of plot to show to the user
            based on simulation results that we have access. \n
            
            Based on the USER_INPUT you should either identify the plot requested by the user
            directly, or decide the correct ones if the USER_INPUT is a bit more generic
            and don't specify any plot. \n
            
            To help you, you have access to a list called AVAILABLE_PLOTS where you have all
            plotting possibilities. Each of the sublists is formated as ['plot_type', 'plot_subtype']
            and represents a single plot. They are not interchangeable, a plot type is always tied to
            it's plot subtype. You must only select the required sublists, never modify or mix them. \n
            
            You must output a JSON object containing a single key 'plots'. The value of this key
            will be a list of lists, where each sublist consists of a plot selection. The format of
            the sublists is ['plot_type', 'plot_subtype'], it must stay in this exact order. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            AVAILABLE_PLOTS: {available_plots} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input","available_plots"],
        )
        
    def get_select_commodity_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an specialist at deciding the correct commodities for each type of plot
            based on the USER_INPUT. \n
            
            You will receive some plot selections in PLOT_SELECTIONS, and you need to select at
            least one commodity for each of them to be plotted. The available commodities are in
            COMMODITIES. \n
            
            You must output a JSON object containing a single key 'commodities'. The value of this key
            will be a list of lists, where each sublist consists of a list of selected commodities (it
            should contain at least one entry). The final list must contain one sublist for each entry
            in PLOT_SELECTIONS. \n
            
            All commodities that you select MUST be in COMMODITIES, you are not allowed to guess any
            of them or modify the way they are written, since the tool cannot plot data for inexistent
            commodities. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            PLOT_SELECTIONS: {plot_selections} \n
            COMMODITIES: {commodities} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input","plot_selections","commodities"],
        )
        
    def get_select_year_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are an specialist at deciding the correct year to plot the selected plots. \n
            
            Based on the USER_INPUT and on the PLOT_SELECTIONS, you should select at least one
            year for that plot to be plotted. The available years are in YEARS. \n
            
            You must output a JSON object containing a single key 'years'. The value of this key
            will be a list of lists, where each sublist consists of a list of selected years (it
            should contain at least one entry). The final list must contain one sublist for each entry
            in PLOT_SELECTIONS. \n
            
            All years that you select MUST be in YEARS, you are not allowed to guess any of them or
            modify the way they are written, since the tool cannot plot data for inexistent years. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            PLOT_SELECTIONS: {plot_selections} \n
            YEARS: {years} \n
            Answer:
            <|eot_id|>
            <|start_header_id|>assistant<|end_header_id|>
            """,
            input_variables=["user_input","plot_selections","years"],
        )
    
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Plotting results')
        user_input = self.state['consolidated_input']
        action_history = self.state['action_history']
        context = self.state['context']
        num_steps = self.state['num_steps']
        num_steps += 1

        select_plot_prompt = self.get_select_plot_prompt_template()
        select_commodity_prompt = self.get_select_commodity_prompt_template()
        select_year_prompt = self.get_select_year_prompt_template()

        select_plot_chain = select_plot_prompt | self.json_model | JsonOutputParser()
        select_commodity_chain = select_commodity_prompt | self.json_model | JsonOutputParser()
        select_year_chain = select_year_prompt | self.json_model | JsonOutputParser()

        runs_dir_path = 'CESM/Runs'
        if self.helper.get_simulation_status() == 'runned':
            simulation = f'{self.mod_model.split("/")[-1][:-5]}-Base'
        else:
            simulation = f'{self.base_model.split("/")[-1][:-5]}-Base'

        db_path = os.path.join(runs_dir_path, simulation, 'db.sqlite')
        conn = sqlite3.connect(db_path)
        dao = DAO(conn)
        plotter = Plotter(dao)

        available_plots = [['Bar', 'ENERGY_CONSUMPTION'],
                            ['Bar', 'ENERGY_PRODUCTION'],
                            ['Bar', 'ACTIVE_CAPACITY'],
                            ['Bar', 'NEW_CAPACITY'],
                            ['Bar', 'CO2_EMISSION'],
                            ['Bar', 'PRIMARY_ENERGY'],
                            ['Sankey', 'SANKEY'],
                            ['SingleValue', 'CAPEX'],
                            ['SingleValue', 'OPEX'],
                            ['SingleValue', 'TOTEX'],
                        #    ['TimeSeries', 'ENERGY_CONSUMPTION'],
                        #    ['TimeSeries','ENERGY_PRODUCTION'],
                        #    ['TimeSeries','POWER_CONSUMPTION'],
                        #    ['TimeSeries','POWER_PRODUCTION'],
                           ]
            
        commodities = [str(c) for c in dao.get_set("commodity")]
        commodities.remove('Dummy')

        years = [int(y) for y in dao.get_set("year")]

        output = select_plot_chain.invoke({"user_input": user_input,
                                            "available_plots": available_plots})
        selected_plots = output['plots']

        final_plots = []
        need_commodity_plots = []
        need_year_plots = []
        for plot in selected_plots:
            if plot[0] == 'Sankey':
                need_year_plots.append(plot + ['NO_COMMODITY'])
            if plot[0] == 'SingleValue' or plot[1] in ['CO2_EMISSION', 'PRIMARY_ENERGY']:
                final_plots.append(plot)
            else:
                need_commodity_plots.append(plot)

        output = select_commodity_chain.invoke({"user_input": user_input,
                                                "plot_selections": need_commodity_plots,
                                                "commodities": commodities})
        selected_commodities = output['commodities']

        for plot, commodities in zip(need_commodity_plots, selected_commodities):
            for commodity in commodities:
                if plot[0] == 'TimeSeries':
                    need_year_plots.append(plot + [commodity])
                else:
                    final_plots.append(plot + [commodity])

        output = select_year_chain.invoke({"user_input": user_input,
                                            "plot_selections": need_year_plots,
                                            "years": years})

        selected_years = output['years']
        for plot, years in zip(need_year_plots, selected_years):
            for year in years:
                if len(plot) == 3:
                    final_plots.append(plot + [year])
                else:
                    final_plots.append(plot + ['', year])
                
        successful_plots = False

        if self.debug:
            self.helper.save_debug('---PLOT RESULTS---')
            self.helper.save_debug(f'SELECTED PLOTS: {final_plots}\n')

        for plot in final_plots:
            plot_type = plot[0]
            plot_subtype = plot[1]
            if len(plot) == 3:
                commodity = plot[2]
                if not(commodity in commodities) and not(commodity == 'NO_COMMODITY'):
                    print(f'{commodity} not in {commodities}')
                    continue
            if len(plot) == 4:
                year = plot[3]
                if not(year in years):
                    print(f'{year} not in {years}')
                    continue
            
            try:
                p_type = getattr(PlotType, plot_type)
                
                if plot_type == 'Bar':
                    # Combination
                    if plot_subtype in  ['PRIMARY_ENERGY', 'CO2_EMISSION']:
                        plotter.plot_bars(getattr(p_type, plot_subtype))
                    else:
                        plotter.plot_bars(getattr(p_type, plot_subtype), commodity=commodity)

                elif plot_type == 'TimeSeries':
                    plotter.plot_timeseries(getattr(p_type, plot_subtype), year=year, commodity=commodity)
                
                elif plot_type == 'Sankey':
                    f = plotter.plot_sankey(year)
                    f.show()

                elif plot_type == 'SingleValue':
                    plotter.plot_single_value([getattr(p_type, plot_subtype)])

                successful_plots = True
            except Exception as e:
                self.helper.save_debug(e)
            
        if successful_plots:
            message = """
            The requested data was successfully plotted and shown to the user, there is no need to try to manipulate any
            data to generate a visual plot to the user. Simply tell him that it was plotted and the data is now available
            for him to check.
            """
        else:
            message = """
            The requested data was not successfully plotted, however, there is no need to try to manipulate any
            data to generate a visual plot to the user. Simply tell him that a problem occured and he should try
            again later.
            """
            
        self.state['num_steps'] = num_steps
        action_history['plot'] = 'done'
        self.state['action_history'] = action_history
        self.state['context'] = context + [message]
        
        return self.state 

# TODO the outputs should also indicate if the model was runned etc...
    
class OutputGenerator(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are part of the Energy System Insight Tool (ESIT), a tool tailored
            to use LLM agents to help users to analyze, manipulate and understand
            energy system models. You are the last agent of the tool, responsible
            for summing up the results generated by other agents in the output
            to be given to the user. \n
            
            Given the USER_INPUT and a CONTEXT, generate an answer for the query
            asked by the user. You should make use of the provided information
            to answer the user in the best possible way. If you think the answer
            does not answer the user completely, ask the user for the necessary
            information if possible. \n
            
            Since you are part of a model analyzing tool you will probably receive
            information about model modifications, model runs, model results, etc.
            If the user is talking about manipulations in the model you should 
            restrain your answer to what was actually done regarding the model
            (information that you'll find in CONTEXT and ACTION_HISTORY), never
            make up modifications or indicate plots for yourself. Other nodes
            already manipulated the model and showed plots to the user, you should
            just forward this information to the user. For listing information
            about parameters you should show a topic list to the user if possible.
            Your main goal in this kind of scenario is to not only tell the user
            what was done to the model, but also, provide as much information as
            possible regarding the model results. Never sum up so much that the
            details about simulation results are lost. \n
            
            CHAT_HISTORY can also be used to gather context and information about
            past messages exchanged between you and the user. \n
            
            Also, you can look at the ACTION_HISTORY to understand which actions were
            executed regarding the model. \n
            
            If and only if the context you use to answer the user provides you a
            data source you should display the provided sources as follows:
            Source:
            - <url/document>
            - <url/document>
            - And so on for how many sources are needed... \n
            
            NEVER output an empty source, if there is no provided source DON'T
            ADD THE SOURCE LIST. THE PRESENCE OF THE SOURCE LIST IS OPTIONAL,
            IF THERE ARE NO SOURCES IN THE DATA YOU USED TO GENERATE THE ANSWER, THEN
            DON'T CREATE THE SOURCE LIST. \n            
            
            It's important never to cite the variables you receive, answer the most
            naturally as possible trying to make it as it was a simple conversation. \n
            
            You will always have access to the current time, but don't mention it unless
            necessary for what the user asked to the tool. \n

            <|eot_id|><|start_header_id|>user<|end_header_id|>
            USER_INPUT: {user_input} \n
            CONTEXT: {context} \n
            ACTION_HISTORY: {action_history} \n
            CHAT_HISTORY: {history} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["user_input","context","action_history","history"],
        )
        
    def execute(self) -> GraphState:
        self.helper.save_chat_status('Generating output')
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.ht_model | StrOutputParser()
        
        ## Get the state
        user_input = self.state['user_input']
        context = self.state['context']
        action_history = self.state['action_history']
        history = self.state['history']
        num_steps = self.state['num_steps']
        num_steps += 1

        llm_output = llm_chain.invoke({"user_input": user_input, "context": context, "action_history": action_history, "history": history})
        
        if self.debug:
            self.helper.save_debug("---GENERATE OUTPUT---")
            self.helper.save_debug(f'GENERATED OUTPUT:\n{llm_output}\n')
        
        if '\nSource:\n- None' in llm_output:
            llm_output = llm_output.replace('\nSource:\n- None','')
            
        self.state['num_steps'] = num_steps
        self.state['final_answer'] = llm_output
        
        return self.state
    
class OutputTranslator(AgentBase):
    def get_prompt_template(self) -> PromptTemplate:
        return PromptTemplate(
            template="""<|begin_of_text|><|start_header_id|>system<|end_header_id|>
            You are the final node of a tool, and you are responsible for translating the
            TOOL_OUTPUT from english to a given TARGET_LANGUAGE. \n
            
            Your output must be a JSON object with a single key, 'output', where you should
            put the translated output. \n
            
            <|eot_id|><|start_header_id|>user<|end_header_id|>
            TOOL_OUTPUT : {tool_output} \n
            TARGET_LANGUAGE : {target_language} \n
            <|eot_id|><|start_header_id|>assistant<|end_header_id|>""",
            input_variables=["tool_output", "target_language"],
        )
    
    def execute(self) -> GraphState:
        final_answer = self.state['final_answer']
        target_language = self.state['target_language']
        num_steps = self.state['num_steps']
        num_steps += 1
        
        if self.debug:
            self.helper.save_debug("---TRANSLATE OUTPUT---")
            self.helper.save_debug(f'TARGET LANGUAGE: {target_language}\n')
        
        prompt = self.get_prompt_template()
        llm_chain = prompt | self.ht_json_model | JsonOutputParser()

        llm_output = llm_chain.invoke({"tool_output": final_answer, "target_language": target_language})
        
        self.state['num_steps'] = num_steps
        self.state['final_answer'] = llm_output['output']
        
        return self.state